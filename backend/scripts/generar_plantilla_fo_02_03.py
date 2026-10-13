"""Genera la plantilla FO 02-03: solo el membrete institucional.

La plantilla que venia del formulario de la auditoria tenia ~100.000 celdas
materializadas (estilos aplicados a filas y columnas enteras). Cargarla y
guardarla costaba ~6 s por exportacion. El cuerpo de las tablas lo escribe el
exportador, asi que la plantilla solo necesita las filas 1-6: este script las
rehace con los mismos estilos (titulo, etiquetas azules, recuadros) y deja el
archivo en unos pocos KB. El exportador redefine igual los anchos de columna,
con lo que la plantilla no los fija.

Uso (desde backend/):

    uv run python scripts/generar_plantilla_fo_02_03.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DESTINO = (
    Path(__file__).resolve().parent.parent
    / "app"
    / "assets"
    / "plantillas"
    / "FO 02-03 Conciliación Bancaria.xlsx"
)

TITULO = "Papeles de Trabajo Conciliación Bancaria"
# El mismo formato contable que traia el formulario para las celdas del codigo.
FORMATO_CONTABLE = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '
FORMATO_FECHA = "yyyy-mm-dd"

AZUL = PatternFill("solid", fgColor="FF00B0F0")
CENTRO = Alignment(horizontal="center", vertical="center")
MEDIO = Side(style="medium")
FINO = Side(style="thin")


def _caja_fina() -> Border:
    return Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _escribir_membrete(hoja, *, por_formula: bool) -> None:
    # El recuadro del logotipo se bordea antes de fusionar: despues las celdas
    # que no son la esquina pasan a MergedCell y openpyxl no deja estilarlas.
    hoja["A1"].border = Border(left=MEDIO, right=MEDIO, top=MEDIO, bottom=MEDIO)
    hoja.merge_cells("A1:B3")
    hoja.merge_cells("C2:H2")

    hoja["I1"] = "FO 02-03"
    hoja["I1"].alignment = CENTRO
    hoja["I1"].border = Border(left=MEDIO, right=MEDIO, top=MEDIO)
    hoja["I1"].number_format = FORMATO_CONTABLE

    hoja["C2"] = TITULO
    hoja["C2"].font = Font(name="Calibri", size=20, bold=True)
    hoja["C2"].alignment = CENTRO
    hoja["C2"].border = Border(left=MEDIO, right=MEDIO)

    hoja["I2"] = "Rev 0"
    hoja["I2"].alignment = CENTRO
    hoja["I2"].border = Border(left=MEDIO, right=MEDIO)

    hoja["I3"] = "Vigencia: 05/12/2023"
    hoja["I3"].alignment = CENTRO
    hoja["I3"].border = Border(left=MEDIO, right=MEDIO, bottom=MEDIO)

    for coordenada, texto in (
        ("A5", "Empresa:"),
        ("H5", "Periodo a conciliar:"),
        ("A6", "Banco:"),
        ("H6", "Numero de Cuenta:"),
    ):
        celda = hoja[coordenada]
        celda.value = texto
        celda.fill = AZUL
        celda.alignment = CENTRO
        celda.border = _caja_fina()

    for coordenada in ("B5", "B6", "I5", "I6"):
        celda = hoja[coordenada]
        celda.alignment = CENTRO
        celda.border = _caja_fina()
    hoja["I5"].number_format = FORMATO_FECHA

    if por_formula:
        # La hoja del extracto repite el membrete por referencia, como en el
        # formulario: empresa, banco, periodo y cuenta se cargan una sola vez.
        hoja["B5"] = "=+'Libro Mayor'!B5"
        hoja["B6"] = "=+'Libro Mayor'!B6"
        hoja["I5"] = "=+'Libro Mayor'!I5"
        hoja["I6"] = "=+'Libro Mayor'!I6"
        hoja["I5"].number_format = FORMATO_FECHA

    hoja.row_dimensions[1].height = 20
    hoja.row_dimensions[2].height = 30
    hoja.row_dimensions[3].height = 20


def main() -> None:
    libro = Workbook()
    mayor = libro.active
    mayor.title = "Libro Mayor"
    movimientos = libro.create_sheet("Movimientos Bancarios")

    _escribir_membrete(mayor, por_formula=False)
    _escribir_membrete(movimientos, por_formula=True)

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    libro.save(DESTINO)
    print(f"Plantilla generada en {DESTINO}")


if __name__ == "__main__":
    main()
