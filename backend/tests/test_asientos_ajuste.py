"""Tests de los asientos de ajuste de partida doble (asientos_ajuste.py).

Lo que se prueba, segun las reglas de negocio (reglas_conciliacion.py):
    - Extracto + gasto bancario / retencion-impuesto -> asiento, Debe contracuenta
      / Haber Banco (sale plata).
    - Extracto + debito automatico (pago directo) -> asiento a revisar contra la
      cuenta de conciliacion de banco.
    - Extracto + acreditacion sin comprobante -> NO asiento (revision manual).
    - Mayor (esta en contabilidad pero no en el banco) -> NO asiento: diferencia
      temporal que se arrastra.
    - Cada asiento balancea (Debe = Haber) y la hoja viaja en el .xlsx.
"""

import io
from datetime import date

import openpyxl

from app.core.config import settings
from app.models.schemas import (
    EncabezadoConciliacion,
    MovimientoBancoExport,
    PendienteExport,
    SolicitudExportacion,
)
from app.services import asientos_ajuste, exportador
from app.services.reglas_conciliacion import (
    ACCION_REVISAR,
    TIPO_DEBITO_AUTOMATICO,
    TIPO_PENDIENTE_TEMPORAL,
)


def _banco(**campos) -> PendienteExport:
    base = {"origen": "banco", "fecha": date(2026, 7, 5), "concepto": "COMISION", "debe": 150.0}
    base.update(campos)
    return PendienteExport(**base)


def _mayor(**campos) -> PendienteExport:
    base = {"origen": "pendiente_mayor", "fecha": date(2026, 7, 4), "concepto": "ALQUILER", "debe": 300.0}
    base.update(campos)
    return PendienteExport(**base)


# ----------------------------------------------------------------------
# Construccion de asientos
# ----------------------------------------------------------------------


def test_solo_el_banco_genera_asiento_y_balancea():
    # La comision del banco lleva asiento; la partida del mayor no.
    asientos = asientos_ajuste.construir_asientos(
        pendientes_banco=[_banco()],
        pendientes_xubio=[_mayor()],
    )
    assert len(asientos) == 1
    assert all(a.balancea for a in asientos)
    assert asientos[0].origen == asientos_ajuste.ORIGEN_BANCO
    assert [a.numero for a in asientos] == [1]


def test_gasto_bancario_usa_la_subcategoria_como_contracuenta():
    asiento = asientos_ajuste.asiento_desde_pendiente(1, _banco(concepto="COMISION", debe=150.0))
    assert asiento is not None
    # Sale plata: Debe la contracuenta (gasto) / Haber Banco.
    cuentas = {(l.cuenta, round(l.debe, 2), round(l.haber, 2)) for l in asiento.lineas}
    assert ("Comisiones", 150.0, 0.0) in cuentas
    assert (settings.CUENTA_BANCO, 0.0, 150.0) in cuentas


def test_gasto_bancario_sin_subcategoria_usa_gastos_bancarios():
    asiento = asientos_ajuste.asiento_desde_pendiente(1, _banco(concepto="INTERESES", debe=50.0))
    assert asiento is not None
    cuentas = {l.cuenta for l in asiento.lineas}
    assert settings.CUENTA_GASTOS_BANCARIOS in cuentas


def test_retencion_impuesto_usa_la_cuenta_de_retenciones():
    asiento = asientos_ajuste.asiento_desde_pendiente(1, _banco(concepto="SIRCREB", debe=120.0))
    assert asiento is not None
    cuentas = {l.cuenta for l in asiento.lineas}
    assert settings.CUENTA_RETENCIONES in cuentas
    # Sigue saliendo plata: Haber Banco.
    assert (settings.CUENTA_BANCO, 0.0, 120.0) in {
        (l.cuenta, round(l.debe, 2), round(l.haber, 2)) for l in asiento.lineas
    }


def test_credito_en_banco_no_genera_asiento():
    # Una acreditacion sin comprobante es un ingreso no registrado: no se arma
    # asiento automatico (la contracuenta es la cuenta del cliente).
    asiento = asientos_ajuste.asiento_desde_pendiente(
        1, _banco(concepto="TRANSFERENCIA", debe=0.0, haber=800.0)
    )
    assert asiento is None

    clasificacion = asientos_ajuste.clasificar_pendiente(
        _banco(concepto="TRANSFERENCIA", debe=0.0, haber=800.0)
    )
    assert clasificacion.accion == ACCION_REVISAR


def test_debito_automatico_genera_asiento_a_revisar():
    # Un pago directo que no es gasto ni impuesto: se propone asiento contra la
    # cuenta de conciliacion de banco, marcado para revisar la contracuenta.
    asiento = asientos_ajuste.asiento_desde_pendiente(
        1, _banco(concepto="PAGO PROVEEDOR", debe=500.0)
    )
    assert asiento is not None
    assert asiento.tipo == TIPO_DEBITO_AUTOMATICO
    assert asiento.revisar_manual is True
    cuentas = {l.cuenta for l in asiento.lineas}
    assert settings.CUENTA_CONCILIACION_BANCO in cuentas
    assert settings.CUENTA_BANCO in cuentas


def test_partida_del_mayor_no_genera_asiento():
    asiento = asientos_ajuste.asiento_desde_pendiente(1, _mayor(concepto="ALQUILER", debe=300.0))
    assert asiento is None

    partida = asientos_ajuste.partida_temporal(_mayor(concepto="ALQUILER", debe=300.0))
    assert partida.origen == asientos_ajuste.ORIGEN_MAYOR
    assert partida.tipo == TIPO_PENDIENTE_TEMPORAL


def test_partida_sin_importe_no_genera_asiento():
    assert asientos_ajuste.asiento_desde_pendiente(1, _banco(debe=0.0, haber=0.0)) is None


def test_totales_debe_y_haber_coinciden():
    asientos = asientos_ajuste.construir_asientos(
        pendientes_banco=[_banco(), _banco(concepto="IVA", debe=200.0)],
        pendientes_xubio=[_mayor()],
    )
    debe, haber = asientos_ajuste.totales(asientos)
    # 150 (comision) + 200 (IVA); el mayor de 300 no lleva asiento.
    assert debe == haber == 350.0


def test_proponer_asientos_separa_las_partidas_sin_asiento():
    asientos, partidas = asientos_ajuste.proponer_asientos(
        pendientes_banco=[_banco(), _banco(concepto="TRANSFERENCIA", debe=0.0, haber=900.0)],
        pendientes_xubio=[_mayor()],
    )
    assert [a.concepto for a in asientos] == ["COMISION"]
    assert {p.concepto for p in partidas} == {"TRANSFERENCIA", "ALQUILER"}


# ----------------------------------------------------------------------
# Hoja del papel de trabajo
# ----------------------------------------------------------------------


def _solicitud(**cambios) -> SolicitudExportacion:
    base = {
        "encabezado": EncabezadoConciliacion(empresa="ACME S.A.", banco="SANT", periodo="2026-07"),
        "pendientes_banco": [_banco(), _banco(fecha=date(2026, 7, 6), concepto="IVA", debe=200.0)],
        "pendientes_xubio": [_mayor()],
        "movimientos_banco": [
            MovimientoBancoExport(fecha=date(2026, 7, 5), concepto="COMISION", debe=150.0, saldo=-150.0),
        ],
    }
    base.update(cambios)
    return SolicitudExportacion(**base)


def _hoja_asientos(contenido: bytes):
    libro = openpyxl.load_workbook(io.BytesIO(contenido))
    assert exportador.HOJA_ASIENTOS in libro.sheetnames
    return libro[exportador.HOJA_ASIENTOS]


def test_la_hoja_de_asientos_viaja_en_el_archivo():
    hoja = _hoja_asientos(exportador.exportar(_solicitud()))
    encabezados = [
        hoja.cell(row=exportador.FILA_ENCABEZADO_ASIENTOS, column=c).value
        for c in range(1, len(exportador.COLUMNAS_ASIENTOS_AJUSTE) + 1)
    ]
    assert encabezados == exportador.COLUMNAS_ASIENTOS_AJUSTE

    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    assert "Comisiones" in valores
    assert "TOTAL GENERAL" in valores
    # Los dos asientos del banco (comision e IVA); el mayor no lleva asiento.
    assert "Total asiento 1" in valores
    assert "Total asiento 2" in valores
    assert "Total asiento 3" not in valores


def test_la_hoja_de_asientos_balancea_a_nivel_total():
    hoja = _hoja_asientos(exportador.exportar(_solicitud()))
    # La fila TOTAL GENERAL trae el mismo numero en Debe (F) y Haber (G).
    fila_total = next(
        f for f in range(1, hoja.max_row + 1)
        if hoja.cell(row=f, column=4).value == "TOTAL GENERAL"
    )
    debe = hoja.cell(row=fila_total, column=6).value
    haber = hoja.cell(row=fila_total, column=7).value
    assert debe == haber == 350.0


def test_la_hoja_de_asientos_se_escribe_sin_pendientes():
    hoja = _hoja_asientos(exportador.exportar(_solicitud(pendientes_banco=[], pendientes_xubio=[])))
    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    assert any(isinstance(v, str) and v.startswith("Sin partidas sin coincidencia") for v in valores)
    assert "TOTAL GENERAL" not in valores
